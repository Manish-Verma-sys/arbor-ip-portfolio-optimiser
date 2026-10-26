"""CSV / XLSX / PDF exports (FR-15). Every export carries data date, score version, assumptions and disclaimer."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from . import APP_NAME, DISCLAIMER
from .scoring import DIM_LABEL

COLUMNS = [
    ("publication_number", "Publication"), ("title", "Title"), ("jurisdiction", "Jur."), ("status", "Status"),
    ("relationship_type", "Relationship"), ("current_owner", "Current owner"), ("strategic_score", "Score"),
    ("data_coverage", "Coverage %"), ("evidence_confidence", "Confidence %"), ("recommendation_label", "Recommendation"),
    ("review_reason_label", "Why review"), ("gate_status", "Checks"), ("gate_codes", "Check details"),
    ("cost_3y", "Fwd cost 3y (USD)"), ("cost_5y", "Fwd cost 5y (USD)"),
    ("next_deadline", "Next deadline"), ("expiry_estimate", "Expiry (est.)"), ("last_status_check", "Status checked"),
] + [(f"rating:{k}", f"{v} (0-5)") for k, v in DIM_LABEL.items()]


def _meta_lines(portfolio, run) -> list[tuple[str, str]]:
    cfg = run.config_snapshot or {}
    return [
        ("Product", APP_NAME),
        ("Portfolio", portfolio.name),
        ("Data mode", portfolio.data_mode.upper() + (" — fictional demo data" if portfolio.data_mode == "synthetic" else "")),
        ("Data as of", run.as_of_date.isoformat()),
        ("Score run", f"{run.id} at {run.run_at.isoformat(timespec='seconds')}Z"),
        ("Scoring version", run.scoring_version),
        ("Scoring config", f"{cfg.get('name')} (v{cfg.get('version')})"),
        ("Weights", ", ".join(f"{DIM_LABEL[k]} {v:g}" for k, v in (cfg.get("weights") or {}).items())),
        ("Coverage threshold", f"{(cfg.get('thresholds') or {}).get('coverage_min')}%"),
        ("Cost assumptions", f"{(cfg.get('cost_assumptions') or {}).get('assumption_version')}; annuity service fee "
                             f"${(cfg.get('cost_assumptions') or {}).get('annuity_service_fee_per_payment')} per payment"),
        ("Data snapshot hash", run.data_snapshot_hash[:16]),
        ("Disclaimer", DISCLAIMER),
    ]


def _cell(r, k):
    if k.startswith("rating:"):
        v = (r.get("ratings") or {}).get(k.split(":", 1)[1])
        return "UNKNOWN" if v is None else v
    v = r.get(k)
    if isinstance(v, list):
        return "; ".join(v)
    return "" if v is None else v


def to_csv(portfolio, run, rows) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for k, v in _meta_lines(portfolio, run):
        w.writerow([f"# {k}", v])
    w.writerow([])
    w.writerow([c[1] for c in COLUMNS])
    for r in rows:
        w.writerow([_cell(r, k) for k, _ in COLUMNS])
    return buf.getvalue()


def to_xlsx(portfolio, run, rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Recommendations"
    head = PatternFill("solid", fgColor="005376")
    ws.append([c[1] for c in COLUMNS])
    for c in ws[1]:
        c.fill = head
        c.font = Font(color="FFFFFF", bold=True)
    for r in rows:
        vals = []
        for k, _ in COLUMNS:
            v = _cell(r, k)
            if k in ("cost_3y", "cost_5y") and v not in ("", None):
                v = float(v)
            vals.append(v)
        ws.append(vals)
    from openpyxl.utils import get_column_letter
    widths = [20, 48, 6, 11, 16, 30, 8, 11, 12, 16, 40, 9, 36, 14, 14, 13, 13, 20] + [14] * 7
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2, min_col=14, max_col=15):
        for c in row:
            c.number_format = '"$"#,##0.00'
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    m = wb.create_sheet("About this export")
    for k, v in _meta_lines(portfolio, run):
        m.append([k, v])
    m.column_dimensions["A"].width = 22
    m.column_dimensions["B"].width = 110
    for c in m["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(portfolio, run, rows, decisions_count: int = 0) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ocean = colors.HexColor("#005376")
    teal = colors.HexColor("#00A39A")
    cream = colors.HexColor("#FFFAED")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"{APP_NAME} — {portfolio.name}")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Times-Roman", textColor=ocean, fontSize=24, alignment=0, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Times-Roman", textColor=ocean, fontSize=15, spaceBefore=10)
    small = ParagraphStyle("s", parent=ss["Normal"], fontSize=8, leading=10)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9, leading=12)
    eyebrow = ParagraphStyle("e", parent=ss["Normal"], fontSize=8, textColor=teal, fontName="Helvetica-Bold")
    s = run.summary or {}
    story = [Paragraph(APP_NAME.upper() + " · MAINTENANCE RECOMMENDATIONS", eyebrow), Paragraph(portfolio.name, h1),
             Paragraph(f"Generated {datetime.utcnow().isoformat(timespec='minutes')}Z · data as of {run.as_of_date} · {run.scoring_version} · run {run.id}", small),
             Spacer(1, 6)]
    br = s.get("by_recommendation") or {}
    kpi = [["Assets", "Families", "Fwd cost 3y", "Fwd cost 5y", "Retain", "Review", "Abandon", "Not in force"],
           [s.get("assets"), s.get("families"),
            f"${float(s.get('forward_cost', {}).get('3', 0)):,.0f}", f"${float(s.get('forward_cost', {}).get('5', 0)):,.0f}",
            br.get("retain", 0), br.get("review", 0), br.get("abandon", 0), br.get("not_in_force", 0)]]
    t = Table(kpi, hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), ocean), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 9), ("BACKGROUND", (0, 1), (-1, 1), cream),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c4c4c4")), ("PADDING", (0, 0), (-1, -1), 5)]))
    story += [t, Paragraph("Recommendation distribution", h2)]
    dist = [[{"retain": "Retain", "review": "Review", "abandon": "Abandon", "not_in_force": "Not in force – review family"}.get(k, k), v]
            for k, v in (s.get("by_recommendation") or {}).items()]
    if dist:
        td = Table([["Category", "Assets"]] + dist, hAlign="LEFT")
        td.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), teal), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c4c4c4"))]))
        story.append(td)
    story.append(Paragraph("Asset recommendations", h2))
    cols = ["publication_number", "title", "jurisdiction", "status", "strategic_score", "data_coverage", "evidence_confidence",
            "recommendation_label", "review_reason_label", "cost_5y", "next_deadline"]
    head = ["Publication", "Title", "Jur.", "Status", "Score", "Cov.%", "Conf.%", "Outcome", "Why review", "Cost 5y", "Next deadline"]
    data = [head]
    for r in rows:
        line = []
        for k in cols:
            v = r.get(k)
            if k == "title":
                v = Paragraph((v or "")[:90], small)
            elif k == "review_reason_label":
                v = Paragraph((v or "").split(" — ")[0], small)
            elif k == "cost_5y":
                v = "unknown" if r.get("cost_unknown") else (f"${float(v):,.0f}" if v else "$0")
            elif v is None:
                v = "—"
            line.append(v)
        data.append(line)
    tbl = Table(data, repeatRows=1, colWidths=[30 * mm, 68 * mm, 10 * mm, 17 * mm, 12 * mm, 12 * mm, 13 * mm, 30 * mm, 38 * mm, 18 * mm, 21 * mm])
    tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), ocean), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                             ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f3f3")]),
                             ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#d0d0d0"))]))
    story.append(tbl)
    story.append(Paragraph("Assumptions & provenance", h2))
    for k, v in _meta_lines(portfolio, run):
        story.append(Paragraph(f"<b>{k}:</b> {v}", body))
    doc.build(story)
    return buf.getvalue()
